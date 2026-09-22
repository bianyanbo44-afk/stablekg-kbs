"""Create the four editable journal submission documents with python-docx."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'submission/neurocomputing'
TITLE='Exact stability certificates for neural temporal knowledge graph completion'
AFFILIATION='Weihai International College, Beijing Jiaotong University, Weihai 264401, China'
AI_DISCLOSURE=('The author defined the study objectives and scope, set the revision priorities, '
    'and retains responsibility for the scientific interpretation and final manuscript. '
    'OpenAI Codex (OpenAI) assisted literature exploration, methodological development, coding, '
    'execution and analysis of the computational experiments, figure scripting, and manuscript '
    'drafting and revision. Quantitative results and data visualizations were produced by the '
    'reproducible computational workflow described in Methods. The author is responsible for '
    'critically reviewing this material and approving the version submitted for publication.')
AI_COVER=('The study received no external funding, and I declare no related competing interests. '
    'I defined the study objectives and scope and directed its revision priorities. '
    'The manuscript describes the supporting role of OpenAI Codex in the research and preparation '
    'workflow; responsibility for the scientific content and final approval rests with me.')

def document(title):
    d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.9);sec.left_margin=sec.right_margin=Inches(.95)
    for name in ('Normal','Title','Heading 1','Heading 2'):
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if 'Theme' in attr:del fonts.attrib[attr]
        for attr in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+attr),'Times New Roman')
    for element in (d.styles.element,d._element):
        for border in element.xpath('.//w:pBdr'):border.getparent().remove(border)
    normal=d.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.space_after=Pt(8);normal.paragraph_format.line_spacing=1.12
    d.styles['Title'].font.size=Pt(18);d.styles['Title'].paragraph_format.space_after=Pt(14)
    d.styles['Heading 1'].font.size=Pt(12);d.styles['Heading 1'].paragraph_format.space_before=Pt(12)
    d.add_paragraph(title,'Title');return d

def main():
    OUT.mkdir(parents=True,exist_ok=True)
    d=document('Title page');d.add_paragraph(TITLE,'Heading 1')
    d.add_paragraph('Research article submitted to Neurocomputing')
    d.add_paragraph('Yanbo Bian','Heading 1');d.add_paragraph(AFFILIATION)
    d.add_paragraph('Corresponding author: Yanbo Bian\nEmail: 24722081@bjtu.edu.cn')
    d.add_paragraph('Programme: Computer Science and Technology')
    d.add_paragraph('Keywords','Heading 1');d.add_paragraph('Temporal knowledge graph; certified robustness; selective prediction; knowledge editing; incremental inference')
    d.add_paragraph('Funding','Heading 1');d.add_paragraph('This research received no external funding.')
    d.add_paragraph('Competing interests','Heading 1');d.add_paragraph('The author declares no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.')
    d.save(OUT/'Title_page.docx')

    d=document('Highlights')
    highlights=[
      'Exact certificates identify answers invariant to shared temporal-window deletion.',
      'A ratio-to-additive reduction makes renormalized evidence exactly certifiable.',
      'Sparse competitor reduction preserves guarantees over the full entity vocabulary.',
      'Certified selection combines calibrated correctness with explicit edit tolerance.',
      'Local maintenance refreshes affected predictions and certificates after edits.']
    assert all(len(h)<=85 for h in highlights)
    for h in highlights:d.add_paragraph(h,'List Bullet')
    d.save(OUT/'Highlights.docx')

    d=document('Declarations');d.add_paragraph(TITLE)
    for heading,body in [
      ('Author contribution','Yanbo Bian: Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, Data curation, Visualization, Writing – original draft, Writing – review and editing.'),
      ('Funding','This research received no external funding.'),
      ('Competing interests','The author declares no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.'),
      ('Data and code availability','Public dataset URLs, source terms, hashes, code, fixed configurations and query-level source data are provided in the Neurocomputing release: https://github.com/bianyanbo44-afk/stablekg-kbs/tree/neurocomputing-v1'),
      ('Use of generative AI',AI_DISCLOSURE)]:
        d.add_paragraph(heading,'Heading 1');d.add_paragraph(body)
    d.save(OUT/'Declarations.docx')

    d=document('Cover letter');d.add_paragraph('22 September 2026\nEditor-in-Chief\nNeurocomputing')
    d.add_paragraph('Dear Editor,')
    d.add_paragraph('Please consider my research article, “'+TITLE+'”, for publication in Neurocomputing. The paper addresses a practical question in neural knowledge systems: how can an answer remain inspectable and maintainable when a coherent period of supporting evidence is withdrawn?')
    d.add_paragraph('The central contribution is an exactly certifiable interface between a frozen neural predictor and an editable temporal memory. Although deleting evidence renormalizes the memory distribution, a ratio-to-additive reduction yields an exact worst-case window-deletion certificate and a constructive counterexample when certification fails. A sparse competitor reduction preserves the full-vocabulary guarantee, while a reverse index supports local maintenance of both answers and certificates.')
    d.add_paragraph('The study evaluates two temporal backbones with three training seeds on ICEWS14 and ICEWS05-15, with confidence, temperature-scaling and ensemble controls on matched predictions. Analytical bounds, random deletion probes and alternative temporal partitions test the role of exact certification. A separate maintenance experiment extends the computational evaluation to a 2.3-million-event GDELT history. The manuscript includes the central figures and results; executable analyses, source data and editable figures accompany the public code release.')
    d.add_paragraph('The work fits Neurocomputing through its connection between neural temporal representations, reliable inference and efficient knowledge maintenance. Its contribution is a precise inference capability with a proved guarantee and measured operating trade-offs. This combination should interest readers developing trustworthy neural reasoning systems and temporal graph applications.')
    d.add_paragraph(AI_COVER)
    d.add_paragraph('Sincerely,\nYanbo Bian\n'+AFFILIATION+'\n24722081@bjtu.edu.cn')
    d.save(OUT/'Cover_letter.docx')
    print('Created Title_page.docx, Highlights.docx, Declarations.docx and Cover_letter.docx')

if __name__=='__main__':main()
